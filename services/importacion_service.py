import csv
import io
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from core.propiedad import obtener_curso_propietario
from models.alumno import Alumno
from models.inscripcion import Inscripcion
from models.usuario import Usuario
from schemas.alumno import (
    AlumnoRespuesta,
    ImportacionAlumnosRespuesta,
    ImportacionError,
    normalizar_documento,
    normalizar_texto,
)
from schemas.inscripcion import InscripcionRespuesta


MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 500

CAMPOS = {
    "nombre": ("nombre", "nombres"),
    "apellido": ("apellido", "apellidos"),
    "documento": (
        "documento",
        "documento_identidad",
        "documento_de_identidad",
        "dni",
    ),
    "fecha_nacimiento": (
        "fecha_nacimiento",
        "fecha_de_nacimiento",
        "nacimiento",
        "fecha_nac",
    ),
}


def _normalizar_encabezado(valor) -> str:
    texto = "" if valor is None else str(valor)
    sin_acentos = "".join(
        caracter
        for caracter in unicodedata.normalize("NFKD", texto)
        if not unicodedata.combining(caracter)
    )

    return sin_acentos.strip().lower().replace(" ", "_")


def _texto(valor) -> str:
    if valor is None:
        return ""

    if isinstance(valor, datetime):
        return valor.date().isoformat()

    if isinstance(valor, date):
        return valor.isoformat()

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def _fila_vacia(valores) -> bool:
    return all(not _texto(valor) for valor in valores)


def _decodificar_csv(contenido: bytes) -> str | None:
    for codificacion in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue

    return None


def _leer_csv(contenido: bytes) -> list[tuple[int, list]] | None:
    texto = _decodificar_csv(contenido)

    if texto is None:
        return None

    lineas = texto.splitlines()
    primera_linea = next(
        (linea for linea in lineas if linea.strip()),
        "",
    )
    delimitador = ";" if primera_linea.count(";") > primera_linea.count(",") else ","

    try:
        lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
        filas = []

        for numero, valores in enumerate(lector, start=1):
            if _fila_vacia(valores):
                continue

            filas.append((numero, valores))

            if len(filas) > MAX_FILAS + 1:
                break
    except csv.Error:
        return None

    return filas


def _leer_xlsx(contenido: bytes) -> list[tuple[int, list]] | None:
    try:
        libro = load_workbook(
            io.BytesIO(contenido),
            read_only=True,
            data_only=True,
        )
    except Exception:
        return None

    try:
        hoja = libro.active
        filas = []

        for numero, valores in enumerate(
            hoja.iter_rows(values_only=True),
            start=1,
        ):
            if _fila_vacia(valores):
                continue

            filas.append((numero, list(valores)))

            if len(filas) > MAX_FILAS + 1:
                break

        return filas
    finally:
        libro.close()


def _leer_planilla(
    nombre_archivo: str,
    contenido: bytes,
) -> tuple[list[tuple[int, list]] | None, str | None]:
    if len(contenido) > MAX_BYTES:
        return None, "archivo_muy_grande"

    extension = Path(nombre_archivo or "").suffix.lower()

    if extension == ".csv":
        filas = _leer_csv(contenido)
    elif extension == ".xlsx":
        filas = _leer_xlsx(contenido)
    else:
        return None, "archivo_formato_invalido"

    if filas is None:
        return None, "archivo_invalido"

    if len(filas) == 0:
        return None, "archivo_vacio"

    return filas, None


def _mapear_columnas(
    encabezados: list,
) -> tuple[dict[str, int] | None, str | None]:
    indices = {}

    for indice, encabezado in enumerate(encabezados):
        normalizado = _normalizar_encabezado(encabezado)
        if normalizado and normalizado not in indices:
            indices[normalizado] = indice

    mapa = {}

    for campo, alias in CAMPOS.items():
        for nombre in alias:
            if nombre in indices:
                mapa[campo] = indices[nombre]
                break

        if campo not in mapa:
            return None, f"Falta la columna obligatoria: {campo}"

    return mapa, None


def _parsear_fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    texto = _texto(valor)

    for formato in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d/%m/%y",
    ):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue

    return None


def _validar_fila(
    fila: dict,
) -> tuple[dict | None, str, str | None]:
    nombre = normalizar_texto(_texto(fila["nombre"])) or ""
    apellido = normalizar_texto(_texto(fila["apellido"])) or ""
    documento = normalizar_documento(_texto(fila["documento"])) or ""
    fecha_nacimiento = _parsear_fecha(fila["fecha_nacimiento"])

    if len(nombre) < 2:
        return None, documento, "El nombre debe tener al menos 2 caracteres"

    if len(apellido) < 2:
        return None, documento, "El apellido debe tener al menos 2 caracteres"

    if len(documento) < 5 or len(documento) > 20:
        return (
            None,
            documento,
            "El documento debe tener entre 5 y 20 caracteres",
        )

    if fecha_nacimiento is None:
        return (
            None,
            documento,
            "La fecha de nacimiento no es válida (usá AAAA-MM-DD o DD/MM/AAAA)",
        )

    if fecha_nacimiento > date.today():
        return (
            None,
            documento,
            "La fecha de nacimiento no puede ser futura",
        )

    return (
        {
            "nombre": nombre,
            "apellido": apellido,
            "documento_identidad": documento,
            "fecha_nacimiento": fecha_nacimiento,
        },
        documento,
        None,
    )


def importar_alumnos(
    db: Session,
    curso_id: int,
    usuario: Usuario,
    nombre_archivo: str,
    contenido: bytes,
) -> tuple[ImportacionAlumnosRespuesta | None, str | None]:
    curso = obtener_curso_propietario(db, curso_id, usuario)

    if curso is None:
        return None, "curso_inexistente"

    if not curso.activo:
        return None, "curso_inactivo"

    filas, error = _leer_planilla(nombre_archivo, contenido)

    if error is not None:
        return None, error

    _, encabezados = filas[0]
    columnas, error = _mapear_columnas(encabezados)

    if error is not None:
        return None, error

    filas_datos = filas[1:]

    if len(filas_datos) > MAX_FILAS:
        return None, "demasiadas_filas"

    if len(filas_datos) == 0:
        return None, "archivo_vacio"

    alumnos: dict[int, AlumnoRespuesta] = {}
    inscripciones: list[InscripcionRespuesta] = []
    errores: list[ImportacionError] = []
    alumnos_creados = 0
    alumnos_inscriptos = 0
    alumnos_omitidos = 0

    for numero, valores in filas_datos:
        fila = {
            campo: valores[indice] if indice < len(valores) else None
            for campo, indice in columnas.items()
        }
        datos, documento, error_fila = _validar_fila(fila)

        if error_fila is not None:
            errores.append(
                ImportacionError(
                    fila=numero,
                    documento=documento or None,
                    mensaje=error_fila,
                )
            )
            continue

        alumno = db.scalar(
            select(Alumno).where(
                Alumno.documento_identidad == datos["documento_identidad"],
            )
        )
        es_nuevo = alumno is None

        if alumno is not None:
            inscripcion_existente = db.scalar(
                select(Inscripcion.id).where(
                    Inscripcion.alumno_id == alumno.id,
                    Inscripcion.curso_id == curso_id,
                )
            )

            if inscripcion_existente is not None:
                alumnos_omitidos += 1
                continue

            if not alumno.activo:
                errores.append(
                    ImportacionError(
                        fila=numero,
                        documento=datos["documento_identidad"],
                        mensaje=(
                            "El alumno está desactivado y no se puede "
                            "inscribir"
                        ),
                    )
                )
                continue

            inscripcion = Inscripcion(
                alumno_id=alumno.id,
                curso_id=curso_id,
            )
        else:
            alumno = Alumno(
                nombre=datos["nombre"],
                apellido=datos["apellido"],
                documento_identidad=datos["documento_identidad"],
                fecha_nacimiento=datos["fecha_nacimiento"],
            )

        try:
            if es_nuevo:
                db.add(alumno)
                db.flush()
                inscripcion = Inscripcion(
                    alumno_id=alumno.id,
                    curso_id=curso_id,
                )

            db.add(inscripcion)
            db.commit()
            db.refresh(alumno)
            db.refresh(inscripcion)
        except IntegrityError:
            db.rollback()
            errores.append(
                ImportacionError(
                    fila=numero,
                    documento=datos["documento_identidad"],
                    mensaje=(
                        "No se pudo importar la fila por un conflicto de datos"
                    ),
                )
            )
            continue

        if alumno.id not in alumnos:
            alumnos[alumno.id] = AlumnoRespuesta.model_validate(alumno)

        inscripciones.append(
            InscripcionRespuesta.model_validate(inscripcion)
        )

        if es_nuevo:
            alumnos_creados += 1

        alumnos_inscriptos += 1

    return (
        ImportacionAlumnosRespuesta(
            total_filas=len(filas_datos),
            alumnos_creados=alumnos_creados,
            alumnos_inscriptos=alumnos_inscriptos,
            alumnos_omitidos=alumnos_omitidos,
            errores=errores,
            alumnos=list(alumnos.values()),
            inscripciones=inscripciones,
        ),
        None,
    )
